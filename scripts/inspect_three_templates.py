"""Read-only inspection; uses bundled openpyxl, never saves a workbook."""
import json
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[2]
b=openpyxl.load_workbook(root/'20260916市场利率调研.xlsx',data_only=True)
s=b['SGD促销']; data={}
data['report']=[dict(row=r,values=[s.cell(r,c).value for c in range(1,4)]) for r in range(1,s.max_row+1)]
data['merges']=[str(m) for m in s.merged_cells.ranges if m.min_col<=3]
data['summary']=[dict(row=r,values=[b['最高报价汇总 '].cell(r,c).value for c in range(2,10)]) for r in range(3,24)]
b=openpyxl.load_workbook(root/'彩虹表_按MarketRateData更新_20260916_17.59.xlsx',data_only=True)
s=b['SGD Promotional Rate'];data['rainbow']={}
for col in range(1,22,3):
    data['rainbow'][str(col)]={'title':s.cell(2,col).value,'rows':[[s.cell(r,c).value for c in range(col,col+3)] for r in range(3,45)],
        'merges':[str(m) for m in s.merged_cells.ranges if m.min_col==col and m.min_row>=3]}
target=root/'market_rate_project/runs/three-bank-20260925/template-inspection.json'
target.write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str),encoding='utf8')
print(json.dumps({'selected': [r for r in data['report'] if r['values'][0] in ['RHB','CIMB','HLF']], 'summary':data['summary'],
                  'groups':[(k,v['title']) for k,v in data['rainbow'].items()]},ensure_ascii=False,default=str))
