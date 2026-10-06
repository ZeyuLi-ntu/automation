"""Saved-file acceptance for the BOC references and Maybank section correction."""
import argparse
from pathlib import Path
import openpyxl
from market_rates.common import load,save
from market_rates.table_validation import sha


def check(output,sgd_output):
    out=Path(output);p=load(out/'table-validation-plan.json');sp=load(Path(sgd_output)/'table-validation-plan.json')
    bp=load(p['base_plan']);sheet=bp['input_sheet'];checks=[];counts={}
    def ck(name,ok):
        checks.append(dict(name=name,passed=bool(ok)))
        if not ok:raise AssertionError(name)
    f=openpyxl.load_workbook(p['report_output'],data_only=False)
    v=openpyxl.load_workbook(p['report_output'],data_only=True)
    for cur in ['SGD','USD','CNY','AUD','NZD','CAD','HKD','EUR','GBP']:
        ws=f[cur+'挂牌'];vs=v[cur+'挂牌'];bc=1 if cur in ['SGD','CNY'] else 2
        first=bc+2;last=first if cur in ['USD','CNY'] else 18 if cur=='SGD' else 13
        count=0
        for row in range(1,ws.max_row+1):
            if str(ws.cell(row,bc).value).strip()!='BOC':continue
            merges=[m for m in ws.merged_cells.ranges if m.min_row==row and m.min_col==bc and m.max_col==bc]
            end=merges[0].max_row if merges else row
            for r in range(row,end+1):
                for c in range(first,last+1):
                    value=vs.cell(r,c).value
                    if isinstance(value,(int,float)):
                        formula=ws.cell(r,c).value
                        ck(cur+' '+ws.cell(r,c).coordinate+' current BOC reference',isinstance(formula,str) and formula.startswith('=') and sheet in formula)
                        count+=1
        ck(cur+' has current BOC rates',count>0);counts[cur]=count
    wanted={6:1.85,9:1.80,12:1.85};seen=set()
    for g in sp['groups']:
        if g['bank']!='Maybank':continue
        details=g['details'];tenor=details[0]['tenor_value'];seen.add(tenor)
        ck('Maybank '+str(tenor)+' main rate',abs(v['SGD促销'].cell(g['report_row'],3).value-wanted[tenor]/100)<1e-10)
        ck('Maybank standalone terms '+str(tenor),all(d['product_id']=='maybank-sgd-standalone' and d['amount_min']=='20000' for d in details))
    ck('Maybank exactly 6/9/12 months',seen==set(wanted))
    for kind in ['report','rainbow']:
        ck(kind+' correction predecessor unchanged',sha(sp[kind+'_template'])==sp[kind+'_sha256'])
    rf=openpyxl.load_workbook(p['rainbow_output'],data_only=False)
    rv=openpyxl.load_workbook(p['rainbow_output'],data_only=True)
    for g in sp['rainbow_groups']:
        for b in g['blocks']:
            if not b.get('group') or b['group']['bank']!='Maybank':continue
            ck('Maybank rainbow '+str(g['tenor']),abs(rv['SGD Promotional Rate'].cell(b['target_start'],g['col']+1).value-float(b['group']['main_pct'])/100)<1e-10)
    ck('Rainbow BOC detail sheet present',sheet in rf.sheetnames)
    result=dict(passed=len(checks),total=len(checks),boc_current_formula_cells=counts,checks=checks,
                report_sha256=sha(p['report_output']),rainbow_sha256=sha(p['rainbow_output']))
    save(out/'boc-maybank-correction-checks.json',result)
    print({k:v for k,v in result.items() if k!='checks'})
    return result


if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--output',required=True);a.add_argument('--sgd-output',required=True);x=a.parse_args();check(x.output,x.sgd_output)
