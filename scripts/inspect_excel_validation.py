"""Read-only inspection of supplied workbook templates, including style and OOXML features."""
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
import openpyxl

ROOT = Path(__file__).resolve().parents[2]


def inspect(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    output = {"file": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheets": {}}
    for ws in wb:
        target = ws.title in {"SGD促销", "最高报价汇总 ", "SGD Promotional Rate"}
        detail = {"size": [ws.max_row, ws.max_column], "merges": [str(m) for m in ws.merged_cells.ranges],
                  "formulas": {c.coordinate: c.value for row in ws for c in row if c.data_type == 'f'},
                  "tables": list(ws.tables), "charts": len(ws._charts), "images": len(ws._images),
                  "freeze": str(ws.freeze_panes), "conditional_formats": len(ws.conditional_formatting),
                  "validations": len(ws.data_validations.dataValidation)}
        if target:
            detail["cells"] = {c.coordinate: {"value": str(c.value) if c.is_date else c.value,
                        "cached": str(values[ws.title][c.coordinate].value), "format": c.number_format,
                        "style_id": c.style_id, "fill": str(c.fill.fgColor), "font": str(c.font),
                        "alignment": str(c.alignment)} for row in ws for c in row if c.value is not None}
            detail["rows"] = {str(r): {"height": d.height, "hidden": d.hidden} for r,d in ws.row_dimensions.items()}
            detail["columns"] = {str(c): {"width": d.width, "hidden": d.hidden} for c,d in ws.column_dimensions.items()}
        output["sheets"][ws.title] = detail
    with ZipFile(path) as z:
        output["parts"] = z.namelist()
    return output


if __name__ == '__main__':
    out = ROOT / 'market_rate_project/outputs/01a0d32d-scb-table-validation'
    out.mkdir(parents=True, exist_ok=True)
    for label, name in [('report', '20260916市场利率调研.xlsx'), ('rainbow', '彩虹表_按MarketRateData更新_20260916_17.59.xlsx')]:
        data = inspect(ROOT / name)
        (out / (label+'-template-inspection.json')).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf8')
        print(label, data['sha256'])
        for name, sheet in data['sheets'].items():
            print(name, sheet['size'], 'formulas', len(sheet['formulas']), 'merges',len(sheet['merges']))
            if 'cells' in sheet:
                rows = {int(''.join(filter(str.isdigit, a))) for a,c in sheet['cells'].items() if 'SCB' in str(c['value'])}
                for a, c in sheet['cells'].items():
                    row=int(''.join(filter(str.isdigit,a)))
                    col=''.join(filter(str.isalpha,a))
                    if row <= 4 or any(abs(row-r)<=1 for r in rows):
                        if len(col)==1 and col <= ('U' if label=='rainbow' else 'I'):
                            print(a, str(c['value'])[:350])
