"""Read-only inspection of the actual approved baseline, never a hard-coded row map."""
import argparse
import json
from pathlib import Path
import openpyxl
from market_rates.common import save
from market_rates.weekly_history import date_value

def inspect(report,rainbow,out):
    b=openpyxl.load_workbook(report,data_only=True);s=b['SGD促销'];data={}
    data['report']=[dict(row=r,values=[s.cell(r,c).value for c in range(1,4)]) for r in range(1,s.max_row+1)]
    data['merges']=[str(m) for m in s.merged_cells.ranges if m.min_col<=3]
    data['report_last_col']=s.max_column;data['baseline_date']=date_value(s['C3'].value)
    data['sheet_names']=b.sheetnames
    data['style_row']=next(r for r in range(4,s.max_row+1) if s.cell(r,1).value=='BEA' and not any(m.min_row<=r<=m.max_row and m.min_col<=1<=m.max_col for m in s.merged_cells.ranges))
    data['summary']=[dict(row=r,values=[b['最高报价汇总 '].cell(r,c).value for c in range(2,10)]) for r in range(3,24)]
    b=openpyxl.load_workbook(rainbow,data_only=True);s=b['SGD Promotional Rate'];data['rainbow']={}
    data['rainbow_last_row']=max(r for r in range(3,s.max_row+1) if any(s.cell(r,c).value is not None for c in range(1,s.max_column+1)))
    for col in range(1,s.max_column+1,3):
        if not s.cell(2,col).value:continue
        data['rainbow'][str(col)]=dict(title=s.cell(2,col).value,rows=[[s.cell(r,c).value for c in range(col,col+3)] for r in range(3,data['rainbow_last_row']+1)],merges=[str(m) for m in s.merged_cells.ranges if m.min_col==col and m.min_row>=3])
    save(out,json.loads(json.dumps(data,default=str)))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--report',required=True);p.add_argument('--rainbow',required=True);p.add_argument('--out',required=True);a=p.parse_args();inspect(a.report,a.rainbow,a.out)
