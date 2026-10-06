"""Saved-workbook checks for the September 28 user corrections."""
import argparse
from collections import Counter
from pathlib import Path
from decimal import Decimal
import openpyxl
from market_rates.common import load,save


def check(out):
    out=Path(out);p=load(out/'table-validation-plan.json');checks=[]
    def ck(name,value):
        if not value:raise AssertionError(name)
        checks.append(name)
    ds=p['details']
    for bank,cur,count in [('HLF','SGD',38),('ICBC','SGD',15),('ICBC','USD',10),('ICBC','CNY',10)]:
        rows=[d for d in ds if d['bank']==bank and d['currency']==cur]
        ck(bank+'/'+cur+' complete numeric details',len(rows)==count and all(d['insertable'] for d in rows))
    hsbc=[d for d in ds if d['bank']=='HSBC' and d['source_date'] and d['source_date']>p['as_of'] and d['currency']!='CHF']
    ck('HSBC later quote-date preserved without pretending it is effective-from',len(hsbc)==168 and all(d['valid_from'] is None and d['insertable'] for d in hsbc))
    ck('No missing-p.a. hold',not any('年化' in d['hold_reason'] or '单位' in d['hold_reason'] for d in ds))
    ck('Same day did not insert duplicate history dates',all(not h['insert_date'] for h in p['histories']))
    for kind in ['report','rainbow']:
        w=openpyxl.load_workbook(p[kind+'_output'],data_only=True);s=w[p['input_sheet']]
        ck(kind+' separates source quote date',s['R3'].value=='银行标注日期' and s['L3'].value=='生效日')
        for d in ds:
            row=d['input_row']
            ck(kind+' percent display '+str(row),s.cell(row,10).value==d['rate_pct']+'%')
            if d in hsbc:
                ck(kind+' HSBC dates '+str(row),s.cell(row,18).value==d['source_date'] and s.cell(row,12).value=='未列明' and s.cell(row,13).value=='2026-09-28')
        for bank,cur,tenor,minimum,rate in [('HLF','SGD',1,'10000','0.0005'),('HLF','SGD',3,'500','0.001'),('ICBC','SGD',6,'500','0.012'),('ICBC','USD',1,None,'0.035'),('ICBC','CNY',1,'500','0.011')]:
            rows=[d for d in ds if d['bank']==bank and d['currency']==cur and d['tenor_value']==tenor and d['amount_min']==minimum]
            ck(kind+' source-known '+bank+cur+str(tenor),len(rows)==1 and abs(Decimal(str(s.cell(rows[0]['input_row'],9).value))-Decimal(rate))<Decimal('0.00000001'))
        w.close()
    save(out/'policy-fix-checks.json',dict(passed=True,checks=len(checks),results=checks,held_reasons=dict(Counter(d['hold_reason'] for d in ds if not d['insertable']))))
    print('User-policy saved-file checks:',len(checks),'passed')


if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--output',required=True);check(a.parse_args().output)
