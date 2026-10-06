"""Read-only independent Excel QA for board expansion and preserved history."""
import argparse,json,html,hashlib
from pathlib import Path
from math import isclose
import openpyxl

def check_output(out):
    out=Path(out);p=json.loads((out/'table-validation-plan.json').read_text(encoding='utf8'));checks=[]
    def check(name,yes,detail=None):
        row=dict(name=name,passed=bool(yes))
        if detail is not None and not yes:row['detail']=detail
        checks.append(row)
    def eq(a,b):return isclose(a,b,rel_tol=0,abs_tol=1e-11) if isinstance(a,(float,int)) and isinstance(b,(float,int)) else a==b
    for kind in ['report','rainbow']:
        old=openpyxl.load_workbook(p[kind+'_template'],data_only=False);oldv=openpyxl.load_workbook(p[kind+'_template'],data_only=True)
        w=openpyxl.load_workbook(p[kind+'_output'],data_only=False);v=openpyxl.load_workbook(p[kind+'_output'],data_only=True)
        changed={m['sheet'] for m in p['matrices']}|{h['sheet'] for h in p['histories']}|{'最高报价汇总 '} if kind=='report' else {'SGD Board Rate','USD Rate + Other Currency Rates'}
        for name in old.sheetnames:
            if name not in changed:
                a={c.coordinate:c.value for row in old[name] for c in row if c.value is not None};b={c.coordinate:c.value for row in w[name] for c in row if c.value is not None}
                check(kind+' untouched '+name,a==b)
        for d in p['details']:
            check(kind+' input '+str(d['input_row']),eq(v[p['input_sheet']].cell(d['input_row'],9).value,float(d['rate_pct'])/100 if d['insertable'] else '-'))
        def cellchecks(name,cells):
            for c in cells:check(kind+' '+name+' '+c['address'],eq(v[name][c['address']].value,c['expected']),[v[name][c['address']].value,c['expected']])
        if kind=='report':
            for m in p['matrices']:cellchecks(m['sheet'],m['cells'])
            for name in {m['sheet'] for m in p['matrices']}:
                allowed={c['address'] for m in p['matrices'] if m['sheet']==name for c in m['cells']}
                for m in [m for m in p['matrices'] if m['sheet']==name]:
                    allowed.add(w[name].cell(m['row'],m['amount_col']).coordinate);allowed.add(w[name].cell(m['row'],m['note_col']).coordinate)
                allowed.add('A1' if name=='SGD挂牌' else 'B1')
                diff=[c.coordinate for row in old[name] for c in row if c.coordinate not in allowed and c.value!=w[name][c.coordinate].value]
                check('matrix preserve other banks '+name,not diff,diff[:20])
            for h in p['histories']:
                name=h['sheet'];cc=h['current_col'];offset=int(h['insert']);cellchecks(name,h['cells']);diff=[]
                from openpyxl.worksheet.cell_range import CellRange
                expected_merges={str(CellRange(min_col=m.min_col+int(bool(offset and m.min_col>=cc)),max_col=m.max_col+int(bool(offset and m.max_col>=cc)),min_row=m.min_row,max_row=m.max_row)) for m in old[name].merged_cells}
                check(name+' all original merges preserved',expected_merges==set(map(str,w[name].merged_cells)),sorted(expected_merges-set(map(str,w[name].merged_cells))))
                # All previously published date columns, including blank values.
                maxcol=250 if name=='USD挂牌' else oldv[name].max_column
                for row in oldv[name].iter_rows(min_col=cc if offset else cc+1,max_col=maxcol):
                    for c in row:
                        if not eq(c.value,v[name].cell(c.row,c.column+offset).value):diff.append(c.coordinate)
                check(name+' historical values preserved',not diff,diff[:20])
                check(name+' all new dates',all(v[name].cell(r,cc).value.strftime('%Y-%m-%d')==p['as_of'] for r in h['date_rows']))
                icons=[str(cf.sqref) for cf,rs in w[name].conditional_formatting._cf_rules.items() if any(r.type=='iconSet' for r in rs) and any(a.min_col<=cc<=a.max_col for a in cf.sqref.ranges)]
                check(name+' no quote arrows',not icons,icons)
            cellchecks('最高报价汇总 ',p['summary'])
            check('USD winning quote minimum',v['最高报价汇总 ']['K43'].value==p['usd_minimum'])
            check('CNY winning quote minimum',v['最高报价汇总 ']['J60'].value==p['cny_minimum'])
            for cell in ['B2','B19','B58']:
                check('summary unrelated heading '+cell,w['最高报价汇总 '][cell].value==old['最高报价汇总 '][cell].value)
            # Very small positive source rates must remain positive and visible.
            e=next(m for m in p['matrices'] if m['sheet']=='EUR挂牌' and any(isinstance(c['expected'],float) for c in m['cells']))
            c=next(c for c in e['cells'] if isinstance(c['expected'],float));check('EUR 0.0001 percent precision',eq(v['EUR挂牌'][c['address']].value,.000001) and w['EUR挂牌'][c['address']].number_format=='0.0000%')
        else:
            for block in p['rainbow_blocks']:
                ranks=[g['rank'] for g in block['groups']];check('rank '+block['currency']+block['tenor'],ranks==sorted(ranks,reverse=True))
                check('one bank block '+block['currency']+block['tenor'],len({g['bank'] for g in block['groups']})==len(block['groups']))
                for i,g in enumerate(block['groups']):
                    if i and g['rank']==block['groups'][i-1]['rank']:check('stable tie '+block['currency']+block['tenor']+g['bank'],g['order']>block['groups'][i-1]['order'])
                    for j,row in enumerate(g['rows']):
                        dest=g['target_row']+j;cell=v[block['sheet']].cell(dest,block['col']+1);check('rainbow '+cell.coordinate+' '+block['currency'],eq(cell.value,row['expected']))
                        oldcolor=old[block['sheet']].cell(row['source_row'],block['col']+1).fill.fgColor
                        newcolor=w[block['sheet']].cell(dest,block['col']+1).fill.fgColor
                        check('bank color '+block['currency']+cell.coordinate,oldcolor==newcolor)
            name='USD Rate + Other Currency Rates';cellchecks(name,p['rainbow_cells']);changedcells={c['address'] for c in p['rainbow_cells']}|{'A31','A'+str(66+p['rainbow_shift'])};diff=[]
            for row in old[name]:
                for c in row:
                    if 31<=c.row<=64:continue
                    dest=w[name].cell(c.row+(p['rainbow_shift'] if c.row>=65 else 0),c.column)
                    if dest.coordinate not in changedcells and dest.value!=c.value:diff.append(c.coordinate)
            check('rainbow promo and other untouched cells',not diff,diff[:30])
        errors=[]
        for s in v:
            for row in s:
                for c in row:
                    if c.data_type!='e':continue
                    rr,cc=c.row,c.column
                    if kind=='report':
                        h=next((h for h in p['histories'] if h['sheet']==s.title),None)
                        if h and h['insert'] and cc>h['current_col']:cc-=1
                    elif s.title=='USD Rate + Other Currency Rates' and rr>=65+p['rainbow_shift']:rr-=p['rainbow_shift']
                    if s.title not in oldv.sheetnames or oldv[s.title].cell(rr,cc).data_type!='e':errors.append(s.title+'!'+c.coordinate+' '+str(c.value))
        check(kind+' no new Excel errors',not errors,errors[:30])
        check(kind+' templates unchanged',hashlib.sha256(Path(p[kind+'_template']).read_bytes()).hexdigest()==p[kind+'_sha256'])
    report=dict(passed=all(x['passed'] for x in checks),checks=len(checks),failed=[c for c in checks if not c['passed']],results=checks,human_reviewed=False)
    (out/'roundtrip-checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    lines=['<!doctype html><meta charset="utf-8"><title>挂牌扩展验证</title><style>body{font:16px Microsoft YaHei;max-width:1200px;margin:36px;line-height:1.8}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:7px}summary{cursor:pointer;padding:10px;background:#edf2f7}</style><h1>挂牌扩展验证</h1>',f'<p>{p["as_of"]} · 回读检查 {sum(x["passed"] for x in checks)}/{len(checks)} · 新结果待人工复核</p>']
    lines+=['<p>'+html.escape(t)+'</p>' for t in p['limitations']]
    for k,label in [('report','调研副本'),('rainbow','彩虹表副本')]:lines.append('<p><a href="'+html.escape(Path(p[k+'_output']).name)+'">'+label+'</a></p>')
    lines.append('<p><a href="http://127.0.0.1:8766">打开本地挂牌人工修正入口</a></p>')
    for bank,cur in sorted({(d['bank'],d['currency']) for d in p['details']}):
        rows=[d for d in p['details'] if d['bank']==bank and d['currency']==cur];lines.append(f'<details><summary>{bank} {cur}：{len(rows)} 条</summary><table><tr><th>期限</th><th>利率%</th><th>金额及渠道</th><th>状态</th></tr>')
        for r in rows:
            vals=[str(r['tenor_value'])+r['tenor_unit'],r['rate_pct'],str(r['amount_min'])+' 至 '+str(r['amount_max'])+' '+r['amount_currency']+' / '+r['channel'],'用户已确认' if r['manual_reviewed'] else r.get('hold_reason') or '待复核']
            lines.append('<tr>'+''.join('<td>'+html.escape(x)+'</td>' for x in vals)+'</tr>')
        lines.append('</table></details>')
    (out/'table-validation.html').write_text(''.join(lines),encoding='utf8');print({k:v for k,v in report.items() if k!='results'},flush=True)
    if not report['passed']:raise ValueError('Roundtrip checks failed')
    return report
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',required=True);check_output(p.parse_args().output)
