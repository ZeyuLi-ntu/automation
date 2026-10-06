"""Saved-file QA independent of the Excel writer and collection process."""
import argparse,json,hashlib,html
from pathlib import Path
from math import isclose
import openpyxl
from market_rates.board_coverage import write_audit

def check_output(out):
    out=Path(out);p=json.loads((out/'table-validation-plan.json').read_text(encoding='utf8'));checks=[];preserved_errors=[]
    def ck(name,ok,detail=None):checks.append(dict(name=name,passed=bool(ok),detail=detail if not ok else None))
    def eq(a,b):return isclose(a,b,rel_tol=0,abs_tol=1e-11) if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==b
    def shift(row,inserts):return row+sum(i['count'] for i in inserts if i['before']<=row)
    for kind in ['report','rainbow']:
        w=openpyxl.load_workbook(p[kind+'_output'],data_only=False);v=openpyxl.load_workbook(p[kind+'_output'],data_only=True)
        old=openpyxl.load_workbook(p[kind+'_template'],data_only=False);oldv=openpyxl.load_workbook(p[kind+'_template'],data_only=True)
        ck(kind+' source unchanged',hashlib.sha256(Path(p[kind+'_template']).read_bytes()).hexdigest()==p[kind+'_sha256'])
        ck(kind+' no JPY in active inputs',all(d['currency']!='JPY' for d in p['details']))
        for d in p['details']:ck(kind+' input '+str(d['input_row']),eq(v[p['input_sheet']].cell(d['input_row'],9).value,float(d['rate_pct'])/100 if d['insertable'] else '-'))
        changed={m['sheet'] for m in p['matrices']}|{h['sheet'] for h in p['histories']}|{'最高报价汇总 '} if kind=='report' else {'SGD Board Rate','USD Rate + Other Currency Rates'}
        for name in old.sheetnames:
            if name in changed:continue
            diff=[c.coordinate for row in oldv[name] for c in row if not eq(c.value,v[name][c.coordinate].value)]
            ck(kind+' untouched values '+name,not diff,diff[:20])
        if kind=='report':
            for m in p['matrices']:
                s=v[m['sheet']]
                date_value=s.cell(1,15 if m['currency']=='SGD' else 10).value
                ck(m['sheet']+' current date heading',hasattr(date_value,'strftime') and date_value.strftime('%Y-%m-%d')==p['as_of'])
                for g in m['groups']:
                    ck(m['sheet']+' bank '+g['bank'],s.cell(g['target_row'],m['bank_col']).value==g['label'])
                    for j,r in enumerate(g['rows']):
                        for c in r['values']:ck(m['sheet']+' '+str(g['target_row']+j)+'/'+str(c['col']),eq(s.cell(g['target_row']+j,c['col']).value,c['expected']))
            for h in p['histories']:
                name=h['sheet'];cc=h['current_col'];offset=int(h['insert_date']);diff=[]
                amount_col=openpyxl.utils.get_column_letter(h['bank_col']+1)
                ck(name+' condition column readable',w[name].column_dimensions[amount_col].width>=46)
                # The current delta is recalculated from the updated quote. It
                # is not a historical observation; final cleanup independently
                # verifies every rebuilt delta against its two numeric inputs.
                delta_columns={c.column for row in oldv[name] for c in row if c.value in ['当日增幅','当日最高报价变动']}
                for row in oldv[name].iter_rows(min_col=cc if offset else cc+1,max_col=oldv[name].max_column):
                    for c in row:
                        if c.column in delta_columns:continue
                        if not eq(c.value,v[name].cell(shift(c.row,h['inserts']),c.column+offset).value):diff.append(c.coordinate)
                ck(name+' all historical values',not diff,diff[:20])
                for c in h['existing']:ck(name+' '+c['address'],eq(v[name][c['address']].value,c['expected']))
                for i in h['inserts']:
                    for j,c in enumerate(i['rows']):ck(name+' inserted '+str(i['target_row']+j),eq(v[name].cell(i['target_row']+j,cc).value,c['expected']))
                    ck(name+' bank group '+str(i['target_row']),any(m.min_col==h['bank_col'] and m.max_col==h['bank_col'] and m.min_row==i['target_anchor'] and m.max_row>=i['target_row']+i['count']-1 for m in w[name].merged_cells))
                for row in h['date_rows']:ck(name+' date '+str(row),v[name].cell(row,cc).value.strftime('%Y-%m-%d')==p['as_of'])
                icons=[str(cf.sqref) for cf,rs in w[name].conditional_formatting._cf_rules.items() if any(r.type=='iconSet' for r in rs) and any(a.min_col<=cc<=a.max_col for a in cf.sqref.ranges)]
                ck(name+' no rate arrows',not icons,icons)
            for c in p['summary']:ck('summary '+c['address'],eq(v['最高报价汇总 '][c['address']].value,c['expected']))
        else:
            for b in p['rainbow_blocks']:
                ranks=[g['rank'] for g in b['groups']];ck(b['currency']+b['tenor']+' rank',ranks==sorted(ranks,reverse=True))
                ck(b['currency']+b['tenor']+' unique bank blocks',len({g['bank'] for g in b['groups']})==len(b['groups']))
                for i,g in enumerate(b['groups']):
                    if i and g['rank']==b['groups'][i-1]['rank']:ck('tie '+b['currency']+b['tenor']+g['bank'],g['order']>b['groups'][i-1]['order'])
                    for j,r in enumerate(g['rows']):
                        addr=v[b['sheet']].cell(g['target_row']+j,b['col']+1).coordinate
                        ck('rainbow '+b['currency']+addr,eq(v[b['sheet']][addr].value,r['expected']))
                        ck('color '+b['currency']+addr,w[b['sheet']][addr].fill.fgColor==old[b['sheet']].cell(r['source_row'],b['col']+1).fill.fgColor)
            name='USD Rate + Other Currency Rates'
            for c in p['rainbow_cells']:ck('rainbow fixed '+c['address'],eq(v[name][c['address']].value,c['expected']))
            diff=[c.coordinate for row in oldv[name].iter_rows(max_row=30) for c in row if not eq(c.value,v[name][c.coordinate].value)]
            ck('rainbow promotions untouched',not diff,diff[:20])
            ck('CNH in CNY block',any(c['address'].startswith('J') and '原CNH' in str(c['expected']) for c in p['rainbow_cells']))
        # New errors on active output areas must never ship.
        errors=[]
        for name in changed|{p['input_sheet']}:
            for row in v[name]:
                for c in row:
                    if c.data_type=='e':
                        rr,cc=c.row,c.column
                        h=next((h for h in p['histories'] if h['sheet']==name),None) if kind=='report' else None
                        if h:
                            oldrow=next((r for r in range(1,oldv[name].max_row+1) if shift(r,h['inserts'])==rr),None)
                            rr=oldrow or rr
                            if h['insert_date'] and cc>h['current_col']:cc-=1
                        if name in oldv.sheetnames and oldv[name].cell(rr,cc).data_type=='e' and oldv[name].cell(rr,cc).value==c.value:
                            preserved_errors.append(dict(workbook=kind,sheet=name,cell=c.coordinate,error=c.value,original_cell=oldv[name].cell(rr,cc).coordinate));continue
                        errors.append(name+'!'+c.coordinate+' '+str(c.value))
        ck(kind+' no new formula errors',not errors,errors[:20])
    result=dict(passed=all(c['passed'] for c in checks),checks=len(checks),failed=[c for c in checks if not c['passed']],preserved_template_errors=preserved_errors,results=checks)
    (out/'roundtrip-checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    notes_path=out/'board-coverage-investigations.json'
    coverage=write_audit(p,out,investigations=json.loads(notes_path.read_text(encoding='utf8')) if notes_path.exists() else None)
    result['coverage_summary']={k:coverage[k] for k in ['required_count','pairs_with_numeric_output','missing_or_fully_held','required_pair_availability_complete']}
    # Numeric QA and scope completeness are separate results. A partial review
    # copy is still allowed, but must never be described as all banks complete.
    (out/'roundtrip-checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    lines=['<!doctype html><meta charset="utf-8"><title>挂牌批量检查</title><style>body{font:16px Microsoft YaHei;max-width:1200px;margin:36px;line-height:1.7}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:8px}summary{padding:12px;background:#eef4f8;cursor:pointer}.warn{color:#a35400}</style><h1>挂牌批量检查</h1>',f'<p>{p["as_of"]} · 当前明细 {len(p["details"])} 条 · 数字回读 {sum(c["passed"] for c in checks)}/{len(checks)}</p>',
        f'<p class="warn"><b>挂牌覆盖：Bank List 应有 {coverage["required_count"]} 个银行/币种组合；当前有数值 {coverage["pairs_with_numeric_output"]} 个，缺数值 {coverage["missing_or_fully_held"]} 个。</b>数字检查通过不代表覆盖齐全。<a href="board-coverage.html">打开完整缺失矩阵及原因</a></p>']
    lines += ['<p>'+html.escape(x)+'</p>' for x in p['limitations']]
    if preserved_errors:lines.append('<p class="warn">原模板历史区域已有 '+str(len(preserved_errors))+' 处公式错误，本次原样保留，未引入新增错误。位置：'+html.escape('；'.join(e['sheet']+'!'+e['cell']+' '+e['error'] for e in preserved_errors))+'。</p>')
    for k,label in [('report','调研副本'),('rainbow','彩虹表副本')]:lines.append('<p><a href="'+html.escape(Path(p[k+'_output']).name)+'">'+label+'</a></p>')
    lines.append('<p><a href="http://127.0.0.1:8766">挂牌人工修正入口</a></p>')
    for bank,cur in sorted({(d['bank'],d['display_currency']) for d in p['details']}):
        rs=[d for d in p['details'] if (d['bank'],d['display_currency'])==(bank,cur)];ok=sum(d['insertable'] for d in rs)
        lines.append(f'<details><summary>{html.escape(bank)} {cur}：{len(rs)} 条；可插入 {ok} 条</summary><table><tr><th>原币种</th><th>期限</th><th>利率%</th><th>金额及条件</th><th>状态</th></tr>')
        for d in rs:
            vals=[d['currency'],str(d['tenor_value'])+d['tenor_unit'],d['rate_pct'],str(d['amount_min'])+' 至 '+str(d['amount_max'])+' '+d['amount_currency']+'；'+d['conditions'],('未插表：'+d['hold_reason']) if d['hold_reason'] else ('可插表；用户已确认' if d['manual_reviewed'] else '已插表；自动核验通过')]
            lines.append('<tr>'+''.join('<td>'+html.escape(x)+'</td>' for x in vals)+'</tr>')
        lines.append('</table></details>')
    (out/'table-validation.html').write_text(''.join(lines),encoding='utf8');print({k:v for k,v in result.items() if k!='results'},flush=True)
    if not result['passed']:raise ValueError('Saved workbook checks failed')
    return result

if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--output',required=True);check_output(a.parse_args().output)
